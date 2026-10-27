from pathlib import Path
from html import escape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import zipfile, json

def package(result,root='client_packages'):
    from audit_engine import save
    root=Path(root);root.mkdir(exist_ok=True,parents=True)
    folder=root/(result['domain'].replace('.','_')+'_'+result['scanned_utc'][:10].replace('-','')+'_'+result['scanned_utc'][11:19].replace(':',''))
    folder.mkdir(exist_ok=True)
    raw, csv=save(result,folder/'Technical_Evidence')
    findings=result['findings'];domain=result['domain']
    css='body{font:16px system-ui,Segoe UI,Arial;max-width:960px;margin:40px auto;padding:0 22px;color:#15243b;line-height:1.55}h1{color:#153d5b}article{border:1px solid #d6e1e8;border-radius:12px;padding:18px;margin:12px 0}.pill{display:inline-block;background:#eaf1f6;padding:3px 10px;border-radius:20px}small{color:#53657b}table{border-collapse:collapse;width:100%}td,th{border:1px solid #ccd8e1;padding:9px;text-align:left}'
    cards=''.join(f'<article><h3>{escape(x["control"])} <span class="pill">{escape(x["severity"])}</span></h3><p>{escape(x["observation"])}</p><p><b>Recommended action:</b> {escape(x["recommendation"])}</p></article>' for x in findings)
    html=f'<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>Email security assessment - {escape(domain)}</title><style>{css}</style><h1>Email Security Assessment</h1><p><b>Domain:</b> {escape(domain)}<br><b>Assessed (UTC):</b> {escape(result["scanned_utc"])}</p><p class="pill">Preliminary read-only assessment — not a certification</p><h2>What we checked</h2><p>Public DNS records for email authentication and routing. This does not test actual mailbox security or prove that email spoofing was blocked.</p><h2>Findings and next steps</h2>{cards}<h2>Who should act?</h2><p>Your designated IT provider or authorized DNS/email administrator should validate sending services, approve and implement any changes, and verify mail delivery.</p><h2>Limits</h2><p>{escape(result["disclaimer"])}</p><small>Cyber Ready Clinic capstone demonstration. Review findings with a qualified administrator before changing production email settings.</small></html>'
    (folder/'02_Email_Security_Dashboard.html').write_text(html,encoding='utf-8')
    styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='NoteSmall',parent=styles['BodyText'],fontSize=9,leading=13,spaceAfter=9))
    def pdf(filename,title,sections):
        doc=SimpleDocTemplate(str(folder/filename),pagesize=(612,792),leftMargin=48,rightMargin=48,topMargin=46,bottomMargin=46)
        story=[Paragraph(title,styles['Title']),Spacer(1,12),Paragraph('Domain: '+escape(domain)+' | UTC: '+escape(result['scanned_utc']),styles['NoteSmall']),Spacer(1,12)]
        for heading,body in sections:
            story.append(Paragraph(escape(heading),styles['Heading2']))
            story.append(Paragraph(escape(body),styles['BodyText']))
            story.append(Spacer(1,10))
        doc.build(story)
    sections=[('Assessment scope','Read-only public DNS analysis of email-related records. No mailbox access, test emails, or production changes.'),('Why it matters','Email authentication helps receivers distinguish authorized domain mail from some impersonation attempts. It does not eliminate business email compromise.'),('Findings', ' | '.join(x['control']+': '+x['severity']+' - '+x['observation'] for x in findings)),('Next step','Have the authorized email/DNS administrator validate senders and review technical evidence before changing policies.'),('Limitations',result['disclaimer'])]
    pdf('01_Executive_Summary.pdf','Email Security - Executive Summary',sections)
    plan=[(x['control']+' - '+x['severity'],x['recommendation']) for x in findings]
    plan.extend([('Approval','Obtain written approval from the domain owner before modifying DNS or mail service settings.'),('Verification','Test authorized outbound messages and monitor DMARC aggregate reports before increasing enforcement.'),('Rollback','Document original DNS values and TTLs; prepare reversal steps before change windows.')])
    pdf('03_Remediation_Action_Plan.pdf','Email Security - Action Plan',plan)
    pdf('05_Client_Instructions.pdf','How to Use Your Assessment',[('Open your dashboard','Double-click 02_Email_Security_Dashboard.html in File Explorer; it opens in Edge or Chrome. It is a saved report, not a live scan.'),('Read the summary','Open 01_Executive_Summary.pdf for the plain-language overview.'),('Who fixes problems?','Share 03_Remediation_Action_Plan.pdf and 04_Technical_Findings.xlsx with your IT provider. Do not change DNS records without authorization.'),('Reassess','Run the separate CyberReadyAudit.exe application if provided, click Run Assessment, and compare results. It needs internet access for public DNS lookups.'),('Privacy','Reports can reveal domain infrastructure. Share only with appropriate stakeholders.')])
    wb=Workbook()
    sheet=wb.active
    sheet.title='Findings'
    rows=[['Control','Severity','Observation','Recommendation']]+[[x[k] for k in ('control','severity','observation','recommendation')] for x in findings]
    for row in rows: sheet.append(row)
    rawsheet=wb.create_sheet('Raw DNS')
    rawsheet.append(['Type','Lookup status','Record value'])
    for typ,v in result['raw_evidence'].items():
        for item in v['values'] or ['(no records returned)']:
            rawsheet.append([typ,v['status'],item])
    for ws,widths in [(sheet,[18,16,72,72]),(rawsheet,[18,22,95])]:
        ws.freeze_panes='A2'
        ws.auto_filter.ref=ws.dimensions
        for i,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=width
        for cell in ws[1]:
            cell.fill=PatternFill('solid',fgColor='173B58')
            cell.font=Font(bold=True,color='FFFFFF')
        for row in ws.iter_rows(min_row=2):
            for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
    wb.save(folder/'04_Technical_Findings.xlsx')
    zpath=folder.with_suffix('.zip')
    with zipfile.ZipFile(zpath,'w',zipfile.ZIP_DEFLATED) as z:
        for p in folder.rglob('*'):
            if p.is_file():z.write(p,p.relative_to(folder))
    return zpath
